from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

DOCUMENT_TITLE = "VOLTAMONS"
AUTHOR = "Diego Mejias Pena"
INSTITUTE = "Ins Cami de Mar"
PROFESSORS = "Jordi Magraner, Inigo Garcia, Javier Linares"
DATE = datetime.now().strftime("%d/%m/%Y")


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def create_cover_page(doc: Document) -> None:
    for _ in range(8):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(DOCUMENT_TITLE)
    run.bold = True
    run.font.size = Pt(32)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Document Tecnic")
    run.font.size = Pt(18)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    data = [
        ("Autor:", AUTHOR),
        ("Institut:", INSTITUTE),
        ("Professors:", PROFESSORS),
        ("Data:", DATE),
        ("", ""),
    ]

    for i, (label, value) in enumerate(data):
        cell_label = table.cell(i, 0)
        cell_value = table.cell(i, 1)

        if label:
            run = cell_label.paragraphs[0].add_run(label)
            run.bold = True
            cell_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cell_label.width = Inches(1.5)

            cell_value.paragraphs[0].add_run(value)
            cell_value.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cell_value.width = Inches(4)

    doc.add_page_break()


def add_section(doc: Document, title: str) -> None:
    doc.add_heading(title, level=1)


def add_subsection(doc: Document, title: str) -> None:
    doc.add_heading(title, level=2)


def add_paragraphs(doc: Document, text: str) -> None:
    for line in text.split("\n"):
        line = line.strip()
        if line:
            doc.add_paragraph(line)


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.add_run(item)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)


def add_table_of_contents(doc: Document) -> None:
    doc.add_heading("Taula de continguts", level=1)
    add_bullet_list(
        doc,
        [
            "1. Introduccio",
            "2. Stack tecnologic",
            "3. Arquitectura del sistema",
            "4. Decisions tecniques",
            "5. Disseny de la base de dades",
            "6. Desplegament",
        ],
    )
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    create_cover_page(doc)
    add_table_of_contents(doc)

    add_section(doc, "1. Introduccio")
    add_subsection(doc, "1.1 Descripcio del projecte")
    add_paragraphs(
        doc,
        "Voltamons es una llibreria online desenvolupada amb Laravel, Inertia i React. "
        "El projecte cobreix el flux complet d'un comerç electronic: cataleg, cistella, "
        "checkout, comandes i un panell d'administracio. Tambe inclou un modul d'opinions "
        "i la generacio de factures en PDF.",
    )
    add_subsection(doc, "1.2 Objectius")
    add_bullet_list(
        doc,
        [
            "Oferir una experiencia de compra clara i rapida per a clients.",
            "Permetre la gestio d'estoc i cataleg per part d'administradors.",
            "Registrar comandes amb historic i factura descarregable.",
            "Implementar un sistema d'opinions validat per compra real.",
        ],
    )
    add_subsection(doc, "1.3 Abast")
    add_paragraphs(
        doc,
        "El projecte cobreix l'aplicacio web completa amb frontend i backend, "
        "persistencia en base de dades i serveis addicionals (emails i PDFs). "
        "No inclou passarel.la de pagament real ni integracio amb serveis de transport externs.",
    )
    doc.add_page_break()

    add_section(doc, "2. Stack tecnologic")
    add_subsection(doc, "2.1 Backend")
    add_bullet_list(
        doc,
        [
            "Laravel 12 amb PHP 8.2.",
            "Arquitectura MVC amb Eloquent ORM.",
            "Middleware d'autenticacio i control de rol.",
        ],
    )
    add_subsection(doc, "2.2 Frontend")
    add_bullet_list(
        doc,
        [
            "Inertia.js + React 18 per a SPA sense API REST tradicional.",
            "Vite per al bundling i hot reload.",
            "Tailwind CSS per a la maquetacio i components.",
        ],
    )
    add_subsection(doc, "2.3 Base de dades")
    add_bullet_list(
        doc,
        [
            "MySQL com a motor principal.",
            "Migrations per a l'evolucio de l'esquema.",
            "Seeders per a dades de prova (rols, categories, llibres i usuaris).",
        ],
    )
    add_subsection(doc, "2.4 Serveis i llibreries")
    add_bullet_list(
        doc,
        [
            "DomPDF per a factures en PDF.",
            "Laravel Mail per a confirmacio de comandes.",
            "Axios per a crides HTTP al frontend.",
        ],
    )
    doc.add_page_break()

    add_section(doc, "3. Arquitectura del sistema")
    add_subsection(doc, "3.1 Estructura de carpetes")
    add_bullet_list(
        doc,
        [
            "app/Http/Controllers: logica de negoci i rutes.",
            "app/Models: models Eloquent (Book, Order, OrderItem, Opinion, Category, Subcategory, User).",
            "resources/js: components i pagines React amb Inertia.",
            "resources/views: plantilles Blade per a PDF i emails.",
            "database/migrations i database/seeders: esquema i dades inicials.",
        ],
    )
    add_subsection(doc, "3.2 Diagrama de l'arquitectura")
    add_paragraphs(doc, "[Imatge 1: Diagrama general Laravel + Inertia + React]")
    add_subsection(doc, "3.3 Flux de dades")
    add_bullet_list(
        doc,
        [
            "El client solicita una ruta web (ex. /cataleg).",
            "El controlador consulta Eloquent i aplica filtres.",
            "Inertia envia les dades a la pagina React.",
            "React renderitza el UI i pot fer crides puntuals JSON.",
            "Les accions de compra creen comandes en transaccio i actualitzen estoc.",
        ],
    )
    add_subsection(doc, "3.4 Patrons de disseny")
    add_bullet_list(
        doc,
        [
            "MVC (Model-View-Controller) al backend.",
            "Components reutilitzables al frontend.",
            "Validacio amb Form Requests i regles de Laravel.",
        ],
    )
    add_subsection(doc, "3.5 Que es una SPA i per que s'utilitza")
    add_paragraphs(
        doc,
        "Una SPA (Single Page Application) es una aplicacio web que carrega una sola pagina HTML "
        "i actualitza el contingut de manera dinamica sense recarregar tot el navegador. "
        "A Voltamons, Inertia permet mantenir rutes del servidor de Laravel, pero el renderitzat "
        "de les vistes es fa al client amb React. Aixo dona una navegacio mes fluida i redueix "
        "el temps de carrega entre pagines.",
    )
    doc.add_page_break()

    add_section(doc, "4. Decisions tecniques")
    add_subsection(doc, "4.1 Inertia per a SPA sense API completa")
    add_paragraphs(
        doc,
        "Problema: cal una experiencia SPA pero sense duplicar logica en una API REST.\n"
        "Solucio: Inertia connecta Laravel amb React, compartint rutes i dades.\n"
        "Justificacio: simplifica el desenvolupament i redueix capes.\n"
        "Alternatives: API REST + client separat, descartat per complexitat extra.",
    )
    add_subsection(doc, "4.2 Cistella persistent amb localStorage")
    add_paragraphs(
        doc,
        "Problema: permetre afegir productes sense login.\n"
        "Solucio: la cistella es guarda a localStorage i s'actualitza amb esdeveniments.\n"
        "Justificacio: millora l'UX i evita dependencia de sessio.\n"
        "Alternatives: sessio de servidor, descartada per dependencia d'autenticacio.",
    )
    add_subsection(doc, "4.3 Control d'estoc amb transaccions")
    add_paragraphs(
        doc,
        "Problema: evitar compres amb estoc insuficient.\n"
        "Solucio: transaccio amb lockForUpdate i decrement d'estoc per linia.\n"
        "Justificacio: garanteix consistencia en compres simultanies.\n"
        "Alternatives: control a nivell frontend, insuficient per concurrencia.",
    )
    add_subsection(doc, "4.4 Opinions vinculades a compra real")
    add_paragraphs(
        doc,
        "Problema: evitar opinions falses.\n"
        "Solucio: nomes es permet opinar si hi ha un OrderItem pendent.\n"
        "Justificacio: incrementa la qualitat i credibilitat de les opinions.",
    )
    add_subsection(doc, "4.5 Facturacio en PDF")
    add_paragraphs(
        doc,
        "Problema: necessitat de comprovant de compra.\n"
        "Solucio: plantilla Blade i generacio amb DomPDF.\n"
        "Justificacio: format estandard i reutilitzable.",
    )
    add_subsection(doc, "4.6 Axios en lloc d'AJAX tradicional")
    add_paragraphs(
        doc,
        "Axios es un client HTTP basat en promeses per al navegador i Node. "
        "Es va triar perque ofereix una API mes neta que l'XMLHttpRequest tradicional, "
        "converteix JSON automaticament, permet interceptors i facilita la gestio d'errors.\n"
        "Alternatives: AJAX natiu o jQuery.ajax, descartats per ser mes verbosos i menys flexibles.",
    )
    add_paragraphs(doc, "Exemple real al projecte (previsualitzacio rapida del cataleg):")
    add_code_block(
        doc,
        [
            "const response = await axios.get(route('catalog.preview', { slug }));",
            "setPreview(response.data.book);",
        ],
    )
    add_paragraphs(doc, "Exemple equivalent amb XMLHttpRequest (mes verbos i dificil de mantenir):")
    add_code_block(
        doc,
        [
            "const xhr = new XMLHttpRequest();",
            "xhr.open('GET', route('catalog.preview', { slug }));",
            "xhr.onload = () => {",
            "  if (xhr.status === 200) {",
            "    const data = JSON.parse(xhr.responseText);",
            "    setPreview(data.book);",
            "  }",
            "};",
            "xhr.send();",
        ],
    )
    doc.add_page_break()

    add_section(doc, "5. Disseny de la base de dades")
    add_subsection(doc, "5.1 Relacions principals")
    add_paragraphs(
        doc,
        "Taules principals i relacions:\n"
        "- roles (1) -> users (N)\n"
        "- categories (1) -> subcategories (N)\n"
        "- categories (1) -> books (N)\n"
        "- subcategories (1) -> books (N, opcional)\n"
        "- users (1) -> orders (N)\n"
        "- orders (1) -> order_items (N)\n"
        "- books (1) -> order_items (N)\n"
        "- books (1) -> opinions (N)",
    )
    add_subsection(doc, "5.2 Decisions de disseny")
    add_bullet_list(
        doc,
        [
            "Snapshots a order_items (title_snapshot, isbn_snapshot) per conservar dades del moment de compra.",
            "Camp has_to_comment per controlar opinions valides postcompra.",
            "subcategory_id nullable per permetre llibres sense subcategoria assignada.",
            "original_price per aplicar i revertir descomptes globals.",
        ],
    )
    add_subsection(doc, "5.3 Exemple de migracio (order_items)")
    add_code_block(
        doc,
        [
            "Schema::create('order_items', function (Blueprint $table) {",
            "    $table->id();",
            "    $table->foreignId('order_id')->constrained()->cascadeOnDelete();",
            "    $table->foreignId('book_id')->constrained()->restrictOnDelete();",
            "    $table->string('title_snapshot', 180);",
            "    $table->string('isbn_snapshot', 20);",
            "    $table->decimal('unit_price', 10, 2);",
            "    $table->unsignedInteger('quantity');",
            "    $table->decimal('line_total', 10, 2);",
            "    $table->boolean('has_to_comment')->default(false);",
            "    $table->timestamps();",
            "});",
        ],
    )
    add_subsection(doc, "5.4 Exemple de relacions Eloquent")
    add_code_block(
        doc,
        [
            "class Order extends Model {",
            "    public function items(): HasMany {",
            "        return $this->hasMany(OrderItem::class);",
            "    }",
            "}",
        ],
    )
    doc.add_page_break()

    add_section(doc, "6. Desplegament")
    add_subsection(doc, "6.1 Requisits per executar")
    add_bullet_list(
        doc,
        [
            "PHP 8.2+, Composer.",
            "Node.js + npm.",
            "MySQL.",
        ],
    )
    add_subsection(doc, "6.2 Eines de desplegament")
    add_bullet_list(
        doc,
        [
            "Vite per a build frontend.",
            "Artisan per a migracions i seeders.",
        ],
    )
    add_subsection(doc, "6.3 Instruccions de desplegament")
    add_bullet_list(
        doc,
        [
            "Instal.lar dependencies PHP i JS.",
            "Configurar .env amb base de dades MySQL.",
            "Executar migracions i seeders.",
            "Compilar assets frontend.",
        ],
    )
    add_subsection(doc, "6.4 Consideracions de seguretat")
    add_bullet_list(
        doc,
        [
            "Autenticacio Laravel Breeze.",
            "Middleware de rol per a rutes admin.",
            "Validacions de formulari en backend.",
            "Transaccions en operacions critiques.",
        ],
    )
    add_subsection(doc, "6.5 Limitacions conegudes")
    add_bullet_list(
        doc,
        [
            "Pagament simulat (no hi ha passarel.la real).",
            "Enviament sense integracio amb transportista extern.",
        ],
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save("Documentacio_Tecnica_Voltamons.docx")
    print("Document creat!")


if __name__ == "__main__":
    main()
